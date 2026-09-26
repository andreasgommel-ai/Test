"""
Mercer-branded PowerPoint slide engine.

Provides a reusable set of helpers for building decks that follow Mercer's
visual identity (deep "Mercer blue" anchored on the Marsh McLennan family
palette) with consistent layouts, typography and footer treatment.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn


# ---------------------------------------------------------------------------
# Brand palette  (Mercer / Marsh McLennan family)
# ---------------------------------------------------------------------------
class C:
    NAVY        = RGBColor(0x00, 0x1A, 0x4A)   # deep background navy
    BLUE        = RGBColor(0x00, 0x2C, 0x77)   # Catalina Blue - Mercer primary
    BLUE_MID    = RGBColor(0x01, 0x6D, 0x9E)   # medium persian blue
    TEAL        = RGBColor(0x00, 0xA8, 0xC7)   # vivid blue-green accent
    SKY         = RGBColor(0xA7, 0xE2, 0xF0)   # blizzard blue (light accent)
    INK         = RGBColor(0x1A, 0x1A, 0x1A)   # near-black body text
    SLATE       = RGBColor(0x5B, 0x67, 0x70)   # secondary grey text
    LIGHT_BG    = RGBColor(0xF2, 0xF4, 0xF7)   # light panel background
    WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
    LINE        = RGBColor(0xD6, 0xDC, 0xE3)   # hairline divider

FONT_HEAD = "Georgia"          # serif for headlines (premium feel)
FONT_BODY = "Calibri"          # clean sans for body / data

# 16:9 canvas
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------
def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    _fill(sp, color)
    sp.shadow.inherit = False
    return sp


def textbox(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
            anchor=MSO_ANCHOR.TOP):
    """lines = list of dicts: {text, size, color, bold, italic, font, space_after}"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", align)
        if ln.get("space_after") is not None:
            p.space_after = Pt(ln["space_after"])
        if ln.get("space_before") is not None:
            p.space_before = Pt(ln["space_before"])
        run = p.add_run()
        run.text = ln["text"]
        f = run.font
        f.size = Pt(ln.get("size", 18))
        f.bold = ln.get("bold", False)
        f.italic = ln.get("italic", False)
        f.name = ln.get("font", FONT_BODY)
        f.color.rgb = ln.get("color", C.INK)
    return tb


def bullets(slide, x, y, w, h, items, size=16, color=C.INK, gap=10,
            marker_color=C.TEAL):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(0)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        # square bullet marker via run
        m = p.add_run()
        m.text = "▪  "
        m.font.size = Pt(size)
        m.font.color.rgb = marker_color
        m.font.name = FONT_BODY
        r = p.add_run()
        r.text = it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = FONT_BODY
    return tb


# ---------------------------------------------------------------------------
# shared chrome
# ---------------------------------------------------------------------------
def logo(slide, color=C.WHITE, x=Inches(0.55), y=Inches(0.4)):
    """Simple text wordmark 'MERCER' used as a lightweight brand mark."""
    tb = slide.shapes.add_textbox(x, y, Inches(3), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "MERCER"
    r.font.name = FONT_HEAD
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = color
    # tracking / letter spacing
    rPr = r._r.get_or_add_rPr()
    rPr.set("spc", "300")
    return tb


def footer(slide, page_no, total, dark=False):
    txt_color = C.SKY if dark else C.SLATE
    tb = slide.shapes.add_textbox(Inches(0.55), Inches(7.02),
                                  Inches(10), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Mercer  |  The Impact of AI on Careers in Europe"
    r.font.size = Pt(8)
    r.font.name = FONT_BODY
    r.font.color.rgb = txt_color
    # page number right
    pn = slide.shapes.add_textbox(Inches(12.0), Inches(7.02),
                                  Inches(0.9), Inches(0.35))
    pp = pn.text_frame.paragraphs[0]
    pp.alignment = PP_ALIGN.RIGHT
    rr = pp.add_run()
    rr.text = f"{page_no:02d} / {total:02d}"
    rr.font.size = Pt(8)
    rr.font.name = FONT_BODY
    rr.font.color.rgb = txt_color


def section_header(slide, kicker, title, accent=C.TEAL):
    """Standard light-content slide header with kicker + title + rule."""
    rect(slide, Inches(0.55), Inches(0.55), Inches(0.18), Inches(0.46), accent)
    textbox(slide, Inches(0.85), Inches(0.5), Inches(11.8), Inches(0.4),
            [{"text": kicker.upper(), "size": 11, "color": C.TEAL,
              "bold": True, "font": FONT_BODY}])
    textbox(slide, Inches(0.85), Inches(0.82), Inches(11.8), Inches(0.7),
            [{"text": title, "size": 26, "color": C.BLUE,
              "bold": False, "font": FONT_HEAD}])
    # hairline
    rect(slide, Inches(0.85), Inches(1.5), Inches(11.93), Pt(1.4), C.LINE)
